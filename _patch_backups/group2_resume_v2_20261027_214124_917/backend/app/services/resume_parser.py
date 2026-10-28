"""离线提取 PDF/DOCX 正文及标题分区，不调用 DeepSeek。"""

import io
import re
import zipfile
from pathlib import Path
from typing import Any

import fitz
from docx import Document

MAX_RESUME_TEXT_CHARS = 100_000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024

SECTION_ALIASES = {
    "education": ("教育经历", "教育背景", "教育信息", "education"),
    "experience": ("工作经历", "实习经历", "工作经验", "实习经验", "employment", "experience"),
    "projects": ("项目经历", "项目经验", "项目实践", "项目介绍", "projects", "project experience"),
    "skills": ("专业技能", "技术技能", "个人技能", "技能清单", "技能", "skills", "technical skills"),
    "awards": ("获奖经历", "荣誉奖项", "证书", "荣誉", "awards", "certifications"),
}


def extract_resume_text(content: bytes, extension: str) -> str:
    """提取可选中的文本；图片扫描版 PDF 需要未来加入 OCR。"""
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("文件内容不是有效的 PDF")
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.page_count > 80:
                    raise ValueError("简历最多支持 80 页")
                text = "\n".join(page.get_text(sort=True) for page in document)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 无法解析，可能已损坏或已加密") from exc

    elif extension == ".docx":
        if not content.startswith(b"PK"):
            raise ValueError("文件内容不是有效的 DOCX")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 1000 or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED:
                    raise ValueError("DOCX 解压后体积过大，拒绝处理")
            document = Document(io.BytesIO(content))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("DOCX 无法解析，可能已损坏") from exc

    else:
        raise ValueError("只支持 PDF 和 DOCX 简历")

    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        raise ValueError("未提取到文字，请上传可复制文字的简历；扫描版暂不支持")
    if len(normalized) > MAX_RESUME_TEXT_CHARS:
        raise ValueError("简历文本超过 100000 字符，请精简后重新上传")
    return normalized


def _section_title(line: str) -> str | None:
    candidate = re.sub(r"[\s:：\-—|┃【】\[\]]", "", line).lower()
    for section, names in SECTION_ALIASES.items():
        if candidate in {re.sub(r"\s", "", value).lower() for value in names}:
            return section
    return None


def parse_resume_sections(text: str) -> dict[str, Any]:
    """只从原文中摘录，不补写项目事实，也不推断心理状态。"""
    sections: dict[str, list[str]] = {
        "profile": [], "education": [], "experience": [],
        "projects": [], "skills": [], "awards": [], "other": [],
    }
    current = "profile"
    for line in text.splitlines():
        detected = _section_title(line)
        if detected is not None:
            current = detected
            continue
        sections[current].append(line)

    email_match = re.search(r"[\w.+\-]+@[\w.\-]+\.[A-Za-z]{2,}", text)
    phone_match = re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", text)
    name_match = re.search(r"(?:姓名|Name)\s*[:：]\s*([^\s|,，]{2,30})", text, re.I)

    return {
        "name": name_match.group(1) if name_match else None,
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None,
        "sections": sections,
        "extraction_method": "rule_based_v1",
    }
