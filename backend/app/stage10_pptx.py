"""Stage 10: PPTX text slides -> existing DOCX material indexing pipeline.

Compatibility entrypoint: POST /api/materials/pptx-import.
There is deliberately no separate Chroma indexer or SQLite schema.
"""
from __future__ import annotations

from io import BytesIO
import logging
import posixpath
import re
from pathlib import Path
import os
from uuid import UUID

from app.stage11_local import account_settings
from app.services.presentation_local import original_path
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile

import httpx
from docx import Document
from fastapi import APIRouter, File, HTTPException, Request, Response, UploadFile

from app.stage8b_auth import require_owner_id

router = APIRouter(prefix="/api/materials", tags=["materials-pptx"])
logger = logging.getLogger(__name__)
MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_SLIDES = 150
MAX_XML_BYTES = 8 * 1024 * 1024
MAX_CHARS = 200_000
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_RELS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_SLIDE_PATH = re.compile(r"^ppt/slides/slide\d+\.xml$")


def _slides_in_presentation_order(archive: ZipFile) -> list[str]:
    names = set(archive.namelist())
    if "ppt/presentation.xml" not in names or "ppt/_rels/presentation.xml.rels" not in names:
        raise ValueError("PPTX 缺少演示文稿目录，文件可能已损坏。")

    prs_info = archive.getinfo("ppt/presentation.xml")
    rel_info = archive.getinfo("ppt/_rels/presentation.xml.rels")
    if prs_info.file_size > MAX_XML_BYTES or rel_info.file_size > MAX_XML_BYTES:
        raise ValueError("PPTX 目录文件过大。")
    presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
    relations = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
    mapping: dict[str, str] = {}
    for relation in relations.findall(f"{_RELS}Relationship"):
        target = relation.attrib.get("Target", "")
        if not target or relation.attrib.get("TargetMode") == "External":
            continue
        # OOXML rel targets may be package-root relative or ppt/ relative.
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("ppt", target))
        if _SLIDE_PATH.fullmatch(path):
            mapping[relation.attrib.get("Id", "")] = path

    slide_ids = presentation.findall(f".//{_P}sldId")
    ordered = []
    for sld in slide_ids:
        relid = sld.attrib.get(f"{_R}id", "")
        path = mapping.get(relid)
        if not path or path not in names:
            raise ValueError("PPTX 幻灯片关系缺失，请重新导出文件。")
        ordered.append(path)
    if not ordered:
        raise ValueError("PPTX 中未找到幻灯片。")
    if len(ordered) > MAX_SLIDES:
        raise ValueError(f"PPTX 超过 {MAX_SLIDES} 页，请拆分后上传。")
    return ordered


def extract_pptx_pages(content: bytes) -> list[tuple[int, list[str]]]:
    """Pure-text slide extraction; charts, screenshots and equations are not OCR'ed."""
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("PPTX 大小超过 12 MiB。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            all_infos = archive.infolist()
            if len(all_infos) > 5000 or sum(item.file_size for item in all_infos) > MAX_UNCOMPRESSED:
                raise ValueError("PPTX 内部文件过多或解压后体积过大。")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names:
                raise ValueError("不是合法的 PPTX 文件。")
            ctinfo = archive.getinfo("[Content_Types].xml")
            if ctinfo.file_size > MAX_XML_BYTES:
                raise ValueError("PPTX 文件格式异常。")
            content_types = archive.read("[Content_Types].xml")
            if b"presentationml.presentation.main+xml" not in content_types:
                raise ValueError("不是标准 PPTX，请将 PPT/PPTM 另存为 PPTX。")
            pages = []
            total_chars = 0
            for index, path in enumerate(_slides_in_presentation_order(archive), 1):
                info = archive.getinfo(path)
                if info.file_size > MAX_XML_BYTES:
                    raise ValueError("某页 PPTX 的内部 XML 过大。")
                root = ET.fromstring(archive.read(path))
                paragraphs = []
                for paragraph in root.findall(f".//{_A}p"):
                    value = "".join(element.text or "" for element in paragraph.iter(f"{_A}t")).strip()
                    if value:
                        paragraphs.append(value)
                        total_chars += len(value)
                        if total_chars > MAX_CHARS:
                            raise ValueError("PPTX 总文字超过 200,000 字符，请拆分上传。")
                pages.append((index, paragraphs))
            if not any(lines for _, lines in pages):
                raise ValueError("PPTX 没有可提取的文字（图片型幻灯片需要 OCR）。")
            return pages
    except (BadZipFile, ET.ParseError, KeyError, RuntimeError) as exc:
        raise ValueError("PPTX 格式损坏或无法解析，请在 PowerPoint 中重新保存。") from exc


def convert_pptx_to_docx(content: bytes, original_name: str) -> bytes:
    pages = extract_pptx_pages(content)
    document = Document()
    document.add_paragraph(f"源自 PPTX：{original_name}（仅提取文本，不包含图像/图表）")
    for slide_num, paragraphs in pages:
        document.add_heading(f"第 {slide_num} 页", level=2)
        for line in paragraphs:
            document.add_paragraph(line)
    out = BytesIO()
    document.save(out)
    data = out.getvalue()
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("文字转换后的文档超过 12 MiB，请拆分 PPTX。")
    return data


@router.post("/pptx-import", status_code=201)
async def import_pptx(request: Request, file: UploadFile = File(...)) -> Response:
    # Important: the existing /api/materials route is itself protected by Stage 8B/C.
    # Auth here and on the nested request must both succeed (fail closed).
    require_owner_id(request)
    filename = (file.filename or "").replace("\\", "/").split("/")[-1]
    if not filename.lower().endswith(".pptx"):
        raise HTTPException(415, "仅支持 .pptx（旧版 .ppt 请另存为 .pptx）。")
    if not filename.strip() or len(filename) > 180:
        raise HTTPException(422, "PPTX 文件名无效或过长。")
    try:
        data = await file.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise HTTPException(413, "文件大小不能超过 12 MiB。")
        if not data:
            raise HTTPException(422, "文件为空。")
        try:
            derivative = convert_pptx_to_docx(data, filename)
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
        # Existing material upload and RAG indexer remain the source of truth.
        # Explicitly use a DOCX derivative because the existing importer does
        # not support preserving a native PPTX file in its DB/file schema.
        safe_stem = Path(filename).stem[:80]
        converted_name = f"{safe_stem}_PPTX文字提取版.docx"
        forwarded_headers = {
            key: value for key, value in request.headers.items()
            if key.lower() not in {
                "content-type", "content-length", "host", "connection",
                "transfer-encoding", "accept-encoding", "authorization", "proxy-authorization",
            } and not key.lower().startswith("x-forwarded-")
        }
        # Authorization header may be used by some local deployments.
        if "authorization" in request.headers:
            forwarded_headers["authorization"] = request.headers["authorization"]
        transport = httpx.ASGITransport(app=request.app, raise_app_exceptions=False)
        async with httpx.AsyncClient(transport=transport, base_url="http://127.0.0.1:8000", timeout=None) as client:
            result = await client.post(
                "/api/materials",
                headers=forwarded_headers,
                files={"file": (converted_name, derivative, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        # Store native PPTX only after successful authenticated DOCX indexing.
        # This file stays strictly inside the current account's uploads directory.
        if result.status_code == 201:
            try:
                material_id = str(UUID(result.json()['id']))
                path = original_path(account_settings(require_owner_id(request)).uploads_dir, material_id)
                path.parent.mkdir(parents=True, exist_ok=True)
                temporary = path.with_suffix('.pptx.uploading')
                try:
                    temporary.write_bytes(data)
                    os.replace(temporary, path)
                finally:
                    temporary.unlink(missing_ok=True)
            except Exception as exc:
                logger.exception('原始 PPTX 保存失败，撤销刚创建的资料')
                # The indexed DOCX must not be presented as a successful PPT upload.
                async with httpx.AsyncClient(transport=transport, base_url='http://127.0.0.1:8000', timeout=None) as client:
                    await client.delete(f'/api/materials/{result.json().get("id", "invalid")}', headers=forwarded_headers)
                raise HTTPException(500, 'PPTX 原件未能保存在本机，上传已撤销') from exc
        # Reuse existing DOCX material record and RAG index semantics.
        passthrough = {"content-type": result.headers.get("content-type", "application/json")}
        return Response(content=result.content, status_code=result.status_code, headers=passthrough)
    finally:
        await file.close()
