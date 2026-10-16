"""补充资料 PDF/DOCX 提取、字符切分；无网络请求。"""

import io
import zipfile

import fitz
from docx import Document


MAX_TEXT_CHARS = 200_000
MAX_PAGES = 150
MAX_DOCX_UNCOMPRESSED = 60 * 1024 * 1024


def extract_material_text(
    raw: bytes,
    extension: str
) -> str:

    if extension == ".pdf":

        if not raw.startswith(b"%PDF-"):
            raise ValueError("不是有效的 PDF 文件")

        try:
            with fitz.open(
                stream=raw,
                filetype="pdf"
            ) as doc:

                if doc.page_count > MAX_PAGES:
                    raise ValueError(
                        f"PDF 最多支持 {MAX_PAGES} 页"
                    )

                text = "\n".join(
                    page.get_text(sort=True)
                    for page in doc
                )

        except ValueError:
            raise

        except Exception as exc:
            raise ValueError(
                "PDF 无法解析，可能已损坏或加密"
            ) from exc

    elif extension == ".docx":

        if not raw.startswith(b"PK"):
            raise ValueError("不是有效的 DOCX 文件")

        try:
            with zipfile.ZipFile(
                io.BytesIO(raw)
            ) as archive:

                infos = archive.infolist()

                if (
                    len(infos) > 1000
                    or sum(
                        x.file_size for x in infos
                    ) > MAX_DOCX_UNCOMPRESSED
                ):
                    raise ValueError(
                        "DOCX 解压后的内容过大"
                    )

            doc = Document(io.BytesIO(raw))

            lines = [
                p.text
                for p in doc.paragraphs
            ]

            for table in doc.tables:
                for row in table.rows:
                    lines.append(
                        " | ".join(
                            cell.text
                            for cell in row.cells
                        )
                    )

            text = "\n".join(lines)

        except ValueError:
            raise

        except Exception as exc:
            raise ValueError(
                "DOCX 无法解析，可能已损坏"
            ) from exc

    else:
        raise ValueError("仅支持 PDF 和 DOCX")

    text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    if not text:
        raise ValueError(
            "无法提取文字；扫描版 PDF 暂不支持，"
            "请上传可复制文字的文件"
        )

    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(
            f"资料正文不能超过 {MAX_TEXT_CHARS} 字符"
        )

    return text


def split_text(
    text: str,
    chunk_size: int = 600,
    overlap: int = 100
) -> list[str]:
    """
    按字符滑动窗口切片。

    每块最多 600 字符；
    相邻块重叠 100 字符。

    MVP 阶段采用简单、可预测的实现。
    """

    if not 0 <= overlap < chunk_size:
        raise ValueError(
            "overlap 必须在 0 到 chunk_size-1 之间"
        )

    clean = text.strip()

    if not clean:
        return []

    step = chunk_size - overlap
    chunks = []

    for start in range(0, len(clean), step):

        part = clean[start:start + chunk_size]

        if part.strip():
            chunks.append(part)

        if start + chunk_size >= len(clean):
            break

    return chunks